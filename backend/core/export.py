# backend/core/export.py
# 报告导出：把落库的评估结果 / 照护报告渲染成 Word（.docx）字节流
#
# 为什么用 python-docx 而不是 HTML→PDF：护理机构实际交付物是 Word（要盖章、要手改），
# 而 python-docx 已在依赖里，零额外依赖。

from datetime import datetime
from io import BytesIO
from typing import Any, Optional

from docx import Document
from docx.shared import Pt

DISCLAIMER = "本报告由 AI 辅助生成，仅供护理团队参考，最终处置以专业人员判断为准。"


def _add_title(doc: Document, text: str) -> None:
    heading = doc.add_heading(text, level=0)
    for run in heading.runs:
        run.font.size = Pt(20)


def _add_kv_table(doc: Document, rows: list[tuple[str, Any]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = "" if value is None else str(value)


def _add_disclaimer(doc: Document) -> None:
    doc.add_paragraph()
    para = doc.add_paragraph(DISCLAIMER)
    para.runs[0].italic = True


def build_assessment_docx(payload: dict) -> bytes:
    """评估报告 → Word。

    入参结构来自 `repositories.get_assessment_result`：
    其中 `summary` 列存的是全量附加数据（structured / issues / summary 三块）。
    """
    summary_payload = payload.get("summary") or {}
    summary = summary_payload.get("summary") or {}
    issues = summary_payload.get("issues") or []
    structured = summary_payload.get("structured") or {}

    doc = Document()
    _add_title(doc, "老年人能力评估报告")

    _add_kv_table(doc, [
        ("评估编号", payload.get("assessment_id", "")),
        ("老人编号", payload.get("elder_id", "")),
        ("综合得分", payload.get("weighted_score", 0)),
        ("评估时间", payload.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("姓名 / 年龄段", f"{structured.get('name', '—')} / {structured.get('age', '—')}"),
        ("照护等级建议", summary.get("care_level_recommendation", structured.get("care_level", "—"))),
    ])

    doc.add_heading("一、维度评分", level=1)
    dims = payload.get("dimension_scores") or []
    if dims:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for idx, title in enumerate(("评估维度", "得分", "权重")):
            table.rows[0].cells[idx].text = title
        for d in dims:  # 兼容 detail/summary 两种形状
            row = table.add_row().cells
            row[0].text = str(d.get("dimension", d.get("name", "")))
            row[1].text = str(d.get("score", ""))
            weight = d.get("weight")
            row[2].text = f"{int(weight * 100)}%" if isinstance(weight, (int, float)) else str(weight or "")
    else:
        doc.add_paragraph("（无维度评分数据）")

    doc.add_heading("二、问题清单", level=1)
    if issues:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for idx, title in enumerate(("优先级", "问题描述", "建议措施")):
            table.rows[0].cells[idx].text = title
        for issue in issues:
            row = table.add_row().cells
            row[0].text = str(issue.get("priority", ""))
            row[1].text = str(issue.get("description", ""))
            row[2].text = str(issue.get("suggestion", ""))
    else:
        doc.add_paragraph("（无高风险问题）")

    doc.add_heading("三、整体评价", level=1)
    overall = summary.get("overall_comment") or summary.get("summary") or "（无整体评价）"
    doc.add_paragraph(str(overall))

    _add_disclaimer(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_care_flow_docx(payload: dict) -> bytes:
    """照护流程报告 → Word。

    入参结构来自 `repositories.get_care_flow_report`，`report` 字段是报告正文（字典或字符串）。
    """
    report = payload.get("report") or {}

    doc = Document()
    _add_title(doc, "照护流程报告")

    _add_kv_table(doc, [
        ("会话编号", payload.get("session_id", "")),
        ("老人编号", payload.get("elder_id", "")),
        ("照护计划", payload.get("care_plan_id", "")),
        ("到达阶段", payload.get("stage_reached", "")),
        ("总轮数", payload.get("total_turns", 0)),
        ("生成时间", payload.get("created_at") or datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ])

    doc.add_heading("一、报告正文", level=1)
    if isinstance(report, dict):
        content = (report.get("content") or report.get("summary")
                   or report.get("report") or report.get("overall_comment"))
        if content:
            doc.add_paragraph(str(content))
        # 其余键值对逐项列出，避免遗漏结构化字段
        extras = {k: v for k, v in report.items() if k not in ("content", "summary", "report")}
        if extras:
            doc.add_heading("二、明细", level=1)
            _add_kv_table(doc, [(str(k), v) for k, v in extras.items()])
    else:
        doc.add_paragraph(str(report) if report else "（无报告内容）")

    _add_disclaimer(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_response_headers(filename: str, ascii_name: str = "report.docx") -> dict:
    """下载 Word 的响应头。

    注意：HTTP 头只能放 latin-1 字符，中文文件名必须走 RFC 5987 的 `filename*=UTF-8''…`，
    同时给一个 ASCII 的 `filename=` 兜底（老客户端/命令行工具用它）。
    """
    from urllib.parse import quote

    return {
        "Content-Disposition": (
            f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
        )
    }
